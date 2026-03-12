"""
Excel eksport:
  - Format: No | Ism-Familya | Telefon Raqam | Manzil | Kasb turi | Savdo Turi | Izoh
  - Savdo Turi katakchalari rangli (Ulgurji=yashil, Chakana=sariq, Servis=moviy)
  - Kechki 22:00 da avtomatik kunlik hisobot
  - Qo'lda eksport: mening viloyatim / barcha viloyatlar
"""
import io
import logging
from datetime import datetime, date
from collections import defaultdict

logger = logging.getLogger(__name__)

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from telegram import Update
from telegram.ext import ContextTypes

from config import ADMIN_IDS, REGION_MAP, REGIONS, SAVDO_COLORS
from database import get_user, get_clients, get_today_clients
from keyboards import export_menu_kb, regions_kb

# ─── RANGLAR VA USLUBLAR ──────────────────────────────────────────────────────

_HEADER_FILL = PatternFill("solid", fgColor="BDD7EE")
_HEADER_FONT = Font(bold=True, size=11, color="1F3864")
_ALT_FILL    = PatternFill("solid", fgColor="F2F2F2")
_TOTAL_FILL  = PatternFill("solid", fgColor="1F3864")
_TOTAL_FONT  = Font(bold=True, color="FFFFFF", size=11)
_THIN        = Side(style="thin", color="9DC3E6")
_CELL_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

_SAVDO_FILL    = {k: PatternFill("solid", fgColor=v) for k, v in SAVDO_COLORS.items()}
_DEFAULT_FILL  = PatternFill()

_COL_WIDTHS = [5, 28, 18, 24, 20, 22, 35]
_HEADERS    = ["No", "Ism - Familya", "Telefon Raqam", "Manzil", "Kasb turi", "Savdo Turi", "Izoh"]

# ─── EXCEL YARATUVCHI ─────────────────────────────────────────────────────────

def _apply_header(ws):
    for col, h in enumerate(_HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill      = _HEADER_FILL
        cell.font      = _HEADER_FONT
        cell.border    = _CELL_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"

def _apply_col_widths(ws):
    for col, width in enumerate(_COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col)].width = width

def _write_client_row(ws, row: int, idx: int, c: dict, alt: bool):
    turi   = c.get("savdo_turi") or ""
    values = [idx, c.get("ism") or "", c.get("telefon") or "",
              c.get("manzil") or "", c.get("kasb_turi") or "", turi, c.get("izoh") or ""]
    row_fill = _ALT_FILL if alt else _DEFAULT_FILL
    for col, val in enumerate(values, 1):
        cell           = ws.cell(row=row, column=col, value=val)
        cell.border    = _CELL_BORDER
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        if col == 6 and turi:
            cell.fill      = _SAVDO_FILL.get(turi, row_fill)
            cell.font      = Font(bold=True, color="FFFFFF" if turi == "Servis" else "1F3864")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        elif col == 1:
            cell.fill      = row_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        else:
            cell.fill = row_fill

def _write_legend(ws, start_row: int):
    legend_col = len(_HEADERS) + 2
    ws.cell(row=start_row, column=legend_col, value="Savdo turi ranglari:").font = Font(bold=True)
    for i, (name, color) in enumerate(SAVDO_COLORS.items(), 1):
        cell           = ws.cell(row=start_row + i, column=legend_col, value=name)
        cell.fill      = PatternFill("solid", fgColor=color)
        cell.font      = Font(bold=True, color="FFFFFF" if name == "Servis" else "1F3864")
        cell.border    = _CELL_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(legend_col)].width = 22

def _fill_sheet(ws, clients: list):
    _apply_header(ws)
    _apply_col_widths(ws)
    for i, c in enumerate(clients, 1):
        _write_client_row(ws, row=i + 1, idx=i, c=c, alt=(i % 2 == 0))
    total_row = len(clients) + 2
    ws.cell(row=total_row, column=1, value="Jami:").fill = _TOTAL_FILL
    ws.cell(row=total_row, column=1).font = _TOTAL_FONT
    ws.cell(row=total_row, column=1).alignment = Alignment(horizontal="center")
    cnt_cell = ws.cell(row=total_row, column=2, value=f"{len(clients)} ta mijoz")
    cnt_cell.fill = _TOTAL_FILL
    cnt_cell.font = _TOTAL_FONT
    for col in range(3, len(_HEADERS) + 1):
        ws.cell(row=total_row, column=col).fill = _TOTAL_FILL
    _write_legend(ws, start_row=2)

def build_daily_excel(clients: list, report_date: str) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    ws_all = wb.create_sheet("Umumiy")
    ws_all.sheet_properties.tabColor = "1F3864"
    _fill_sheet(ws_all, clients)
    by_region: dict[int, list] = defaultdict(list)
    for c in clients:
        by_region[c["region_id"]].append(c)
    for region in REGIONS:
        rid = region["id"]
        if rid in by_region:
            ws = wb.create_sheet(region["name"][:31])
            _fill_sheet(ws, by_region[rid])
    return wb

def build_full_excel(clients: list, region_id: int = None) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    if region_id:
        ws = wb.create_sheet(REGION_MAP.get(region_id, "Viloyat")[:31])
        _fill_sheet(ws, clients)
    else:
        ws_all = wb.create_sheet("Umumiy")
        _fill_sheet(ws_all, clients)
        by_region: dict[int, list] = defaultdict(list)
        for c in clients:
            by_region[c["region_id"]].append(c)
        for region in REGIONS:
            rid = region["id"]
            if rid in by_region:
                ws = wb.create_sheet(region["name"][:31])
                _fill_sheet(ws, by_region[rid])
    return wb

# ─── TELEGRAM HANDLERLARI ─────────────────────────────────────────────────────

async def export_menu(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    if update.callback_query:
        await update.callback_query.answer()
    is_admin = update.effective_user.id in ADMIN_IDS
    await update.effective_message.reply_text(
        "📤 *Excel eksport*\n\nQaysi ma'lumotlarni eksport qilmoqchisiz?",
        parse_mode="Markdown",
        reply_markup=export_menu_kb(is_admin=is_admin),
    )

async def export_my_region(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    if update.callback_query:
        await update.callback_query.answer()
    db_user   = await get_user(update.effective_user.id)
    region_id = db_user["region_id"] if db_user else None
    await _do_export(update, ctx, region_id)

async def export_all(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    if update.callback_query:
        await update.callback_query.answer()
    if update.effective_user.id not in ADMIN_IDS:
        await update.effective_message.reply_text("❌ Faqat adminlar uchun!")
        return
    await _do_export(update, ctx, region_id=None)

async def export_choose_region(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    if update.callback_query:
        await update.callback_query.answer()
    if update.effective_user.id not in ADMIN_IDS:
        await update.effective_message.reply_text("❌ Faqat adminlar uchun!")
        return
    ctx.user_data["export_choosing"] = True
    await update.effective_message.reply_text(
        "🗺 Eksport uchun viloyat tanlang:",
        reply_markup=regions_kb(include_all=True),
    )

async def export_region_selected(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    if not ctx.user_data.get("export_choosing"):
        return
    query = update.callback_query
    await query.answer()
    ctx.user_data.pop("export_choosing", None)
    data      = query.data
    region_id = None if data == "region_all" else int(data.split("_")[1])
    await _do_export(update, ctx, region_id)

async def _do_export(update: Update, ctx: ContextTypes.DEFAULT_TYPE, region_id: int = None):
    is_admin = update.effective_user.id in ADMIN_IDS
    loading  = await update.effective_message.reply_text("⏳ Excel fayl tayyorlanmoqda...")

    clients = await get_clients(region_id, limit=10_000)
    if not clients:
        await loading.edit_text("❌ Eksport qilish uchun mijozlar yo'q.")
        await update.effective_message.reply_text(
            "📤 *Excel eksport*", parse_mode="Markdown",
            reply_markup=export_menu_kb(is_admin=is_admin),
        )
        return

    wb     = build_full_excel([dict(c) for c in clients], region_id)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    region_name = REGION_MAP.get(region_id, "Umumiy") if region_id else "Umumiy"
    today       = date.today().strftime("%Y-%m-%d")
    filename    = f"CRM_{region_name}_{today}.xlsx"

    await update.effective_chat.send_document(
        document=buffer,
        filename=filename,
        caption=(
            f"📊 *{region_name} — CRM bazasi*\n"
            f"👥 Jami: *{len(clients)}* ta mijoz\n"
            f"📅 {datetime.now().strftime('%Y-%m-%d %H:%M')}"
        ),
        parse_mode="Markdown",
    )
    await loading.edit_text("✅ Excel fayl yuborildi!")

# ─── 22:00 AVTOMATIK KUNLIK HISOBOT ──────────────────────────────────────────

async def send_daily_report(app):
    from config import ADMIN_IDS
    today   = date.today().strftime("%Y-%m-%d")
    clients = [dict(c) for c in await get_today_clients(region_id=None)]

    for admin_id in ADMIN_IDS:
        try:
            db_user   = await get_user(admin_id)
            region_id = db_user["region_id"] if db_user else None

            if region_id:
                day_clients = [c for c in clients if c["region_id"] == region_id]
                region_name = REGION_MAP.get(region_id, "Viloyat")
            else:
                day_clients = list(clients)
                region_name = "Barcha viloyatlar"

            if not day_clients:
                await app.bot.send_message(
                    chat_id=admin_id,
                    text=(
                        f"📊 *Kunlik hisobot — {today}*\n"
                        f"🗺 {region_name}\n\n"
                        f"Bugun hech qanday mijoz qo'shilmadi."
                    ),
                    parse_mode="Markdown",
                )
                continue

            wb     = build_daily_excel(day_clients, today)
            buffer = io.BytesIO()
            wb.save(buffer)
            buffer.seek(0)

            await app.bot.send_document(
                chat_id=admin_id,
                document=buffer,
                filename=f"Hisobot_{region_name}_{today}.xlsx",
                caption=(
                    f"📊 *Kunlik hisobot — {today}*\n"
                    f"🗺 {region_name}\n"
                    f"👥 Bugun qo'shildi: *{len(day_clients)}* ta mijoz\n"
                    f"⏰ Vaqt: 22:00"
                ),
                parse_mode="Markdown",
            )
        except Exception as e:
            logger.warning(f"22:00 hisobot xato (admin {admin_id}): {e}")

# ─── VAZIFALAR EXCEL VA BILDIRISHNOMA ─────────────────────────────────────────

_TASK_HEADERS    = ["No", "Sarlavha", "Mijoz", "Muddat", "Tavsif", "Viloyat"]
_TASK_COL_WIDTHS = [5, 30, 25, 20, 35, 20]


def build_task_excel(tasks: list) -> "Workbook":
    wb = Workbook()
    ws = wb.active
    ws.title = "Vazifalar"

    # Header
    for col, h in enumerate(_TASK_HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill      = _HEADER_FILL
        cell.font      = _HEADER_FONT
        cell.border    = _CELL_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"
    for col, width in enumerate(_TASK_COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col)].width = width

    # Rows
    for i, t in enumerate(tasks, 1):
        row_fill = _ALT_FILL if i % 2 == 0 else _DEFAULT_FILL
        values = [
            i,
            t.get("sarlavha") or "",
            t.get("client_name") or "—",
            t.get("muddat") or "—",
            t.get("tavsif") or "—",
            t.get("region_name") or "—",
        ]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=i + 1, column=col, value=val)
            cell.border    = _CELL_BORDER
            cell.fill      = row_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if col == 1:
                cell.alignment = Alignment(horizontal="center", vertical="center")

    # Total row
    total_row = len(tasks) + 2
    tc = ws.cell(row=total_row, column=1, value="Jami:")
    tc.fill = _TOTAL_FILL; tc.font = _TOTAL_FONT
    tc.alignment = Alignment(horizontal="center")
    cc = ws.cell(row=total_row, column=2, value=f"{len(tasks)} ta vazifa")
    cc.fill = _TOTAL_FILL; cc.font = _TOTAL_FONT
    for col in range(3, len(_TASK_HEADERS) + 1):
        ws.cell(row=total_row, column=col).fill = _TOTAL_FILL
    return wb


async def notify_admins_new_task(bot, task_id: int, sarlavha: str, client_name: str,
                                  muddat: str, tavsif: str, region_name: str, added_by: str):
    """Yangi muddatli vazifa qo'shilganda adminlarga Excel bilan darhol xabar beradi."""
    today = date.today().strftime("%Y-%m-%d")
    task_row = [{
        "sarlavha":    sarlavha,
        "client_name": client_name,
        "muddat":      muddat,
        "tavsif":      tavsif,
        "region_name": region_name,
    }]
    wb  = build_task_excel(task_row)
    buf = io.BytesIO()
    wb.save(buf)
    # Bytes sifatida saqlash — har bir admin uchun yangi BytesIO yaratiladi
    wb_bytes = buf.getvalue()

    caption = (
        f"📋 *Yangi eslatma vazifasi qo'shildi!*\n\n"
        f"📝 *{sarlavha}*\n"
        f"🔗 Mijoz: {client_name}\n"
        f"🗓 Muddat: {muddat}\n"
        f"📄 Tavsif: {tavsif or '—'}\n"
        f"🗺 Viloyat: {region_name}\n"
        f"👤 Qo'shdi: {added_by}\n"
        f"ID: `{task_id}`"
    )
    for admin_id in ADMIN_IDS:
        try:
            await bot.send_document(
                chat_id=admin_id,
                document=io.BytesIO(wb_bytes),
                filename=f"Vazifa_{task_id}_{today}.xlsx",
                caption=caption,
                parse_mode="Markdown",
            )
        except Exception as e:
            logger.warning(f"Vazifa bildirishnoma xato (admin {admin_id}): {e}")


# ─── MUDDATI O'TGAN VAZIFALAR ─────────────────────────────────────────────────

async def get_overdue_tasks_and_notify(app):
    from database import get_overdue_tasks, mark_task_notified
    tasks = await get_overdue_tasks()
    if not tasks:
        return
    for t in tasks:
        client_info = f"\n🔗 Mijoz: *{t['client_ism']}*" if t["client_ism"] else ""
        msg_text = (
            f"⏰ *Muddati o'tgan vazifa!*\n\n"
            f"📝 *{t['sarlavha']}*{client_info}\n"
            f"📄 {t['tavsif'] or '—'}\n"
            f"🗓 Muddat: {t['muddat']}\n"
            f"🗺 {REGION_MAP.get(t['region_id'], '?')}\n"
            f"ID: `{t['id']}`\n\n"
            f"_/vazifa {t['id']} — boshqarish uchun_"
        )
        sent = False
        for admin_id in ADMIN_IDS:
            try:
                await app.bot.send_message(
                    chat_id=admin_id,
                    text=msg_text,
                    parse_mode="Markdown",
                )
                sent = True
            except Exception as e:
                logger.warning(f"Eslatma xato (admin {admin_id}): {e}")
        # Hech bo'lmaganda bir adminга yuborilsa, qayta xabar bermaylik
        if sent:
            await mark_task_notified(t["id"])
